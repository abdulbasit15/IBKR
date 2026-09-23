"""Read trades from and write analysis back to the ZTH Trade Tracker workbook."""
from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

# 1-indexed column positions in the "Trade Tracker - Eval" sheet.
COL = {
    "zth": 1, "date": 2, "time": 3, "asset": 4, "direction": 5,
    "entry": 6, "exit": 7, "size": 8, "stop": 9, "target": 10,
    "rr_targeted": 11, "rr_realized": 12, "winloss": 13,
    "pnl": 14, "pnl_pct": 15, "cum_pnl": 16,
    "setup": 17, "notes": 18, "instrument_level": 19,
    "level_screenshot": 20, "trade_screenshot": 21,
    "mistakes": 22, "rules": 23, "market": 24, "emotion": 25,
    "process": 26, "learning": 27,
}
FIRST_DATA_ROW = 2


@dataclass
class Trade:
    zth: object = None
    date: object = None
    time: object = None
    asset: str = ""
    direction: str = ""
    entry: float = None
    exit: float = None
    size: float = None
    stop: float = None
    target: float = None
    rr_targeted: float = None
    rr_realized: float = None
    winloss: str = ""
    pnl: float = None
    pnl_pct: float = None
    cum_pnl: float = None
    setup: str = ""
    notes: str = ""
    instrument_level: str = ""
    level_screenshot: str = ""
    trade_screenshot: str = ""
    mistakes: str = ""
    rules: str = ""
    market: str = ""
    emotion: str = ""
    process: str = ""
    learning: str = ""

    @property
    def timestamp(self) -> datetime | None:
        if self.date is None:
            return None
        d = self.date.date() if isinstance(self.date, datetime) else self.date
        t = self.time if isinstance(self.time, time) else time(0, 0)
        return datetime.combine(d, t)

    def dedupe_key(self) -> tuple:
        d = self.date.date() if isinstance(self.date, datetime) else self.date
        t = self.time.strftime("%H:%M") if isinstance(self.time, time) else ""
        return (str(d), t, (self.asset or "").upper(), (self.direction or "").title(),
                round(float(self.entry), 4) if self.entry is not None else None)


def _cell(ws, row, key):
    return ws.cell(row=row, column=COL[key]).value


def read_trades(workbook: Path | None = None) -> list[Trade]:
    """Return the real (non-placeholder) trades using cached computed values."""
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb[config.TRADE_SHEET]
    trades: list[Trade] = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        d, asset, entry = _cell(ws, r, "date"), _cell(ws, r, "asset"), _cell(ws, r, "entry")
        if d is None or not asset or entry is None:
            continue  # placeholder / blank row
        trades.append(Trade(
            zth=_cell(ws, r, "zth"), date=d, time=_cell(ws, r, "time"),
            asset=str(asset).strip().upper(),
            direction=str(_cell(ws, r, "direction") or "").strip().title(),
            entry=_num(entry), exit=_num(_cell(ws, r, "exit")),
            size=_num(_cell(ws, r, "size")), stop=_num(_cell(ws, r, "stop")),
            target=_num(_cell(ws, r, "target")),
            rr_targeted=_num(_cell(ws, r, "rr_targeted")),
            rr_realized=_num(_cell(ws, r, "rr_realized")),
            winloss=str(_cell(ws, r, "winloss") or "").strip(),
            pnl=_num(_cell(ws, r, "pnl")), pnl_pct=_num(_cell(ws, r, "pnl_pct")),
            cum_pnl=_num(_cell(ws, r, "cum_pnl")),
            setup=str(_cell(ws, r, "setup") or "").strip(),
            notes=str(_cell(ws, r, "notes") or "").strip(),
            instrument_level=str(_cell(ws, r, "instrument_level") or "").strip(),
            level_screenshot=str(_cell(ws, r, "level_screenshot") or "").strip(),
            trade_screenshot=str(_cell(ws, r, "trade_screenshot") or "").strip(),
            mistakes=str(_cell(ws, r, "mistakes") or "").strip(),
            rules=str(_cell(ws, r, "rules") or "").strip(),
            market=str(_cell(ws, r, "market") or "").strip(),
            emotion=str(_cell(ws, r, "emotion") or "").strip(),
            process=str(_cell(ws, r, "process") or "").strip(),
            learning=str(_cell(ws, r, "learning") or "").strip(),
        ))
    wb.close()
    return trades


def _num(v):
    if v is None or v == "" or isinstance(v, str):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def backup(workbook: Path | None = None) -> Path:
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    config.BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = config.BACKUP_DIR / f"{wb_path.stem}.{stamp}.bak.xlsx"
    shutil.copy2(wb_path, dest)
    return dest


def write_setups(labels: dict, workbook: Path | None = None,
                 only_blank: bool = True, column: str = "setup",
                 replace_values: set | None = None) -> int:
    """Write labels into a tag column for real trade rows (in read order).

    `labels` is keyed by the 0-based index of each real trade row (matching
    read_trades order). Values may be a string or a list of tags (joined with
    ", "). Returns the number of cells written.
    """
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    backup(wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws = wb[config.TRADE_SHEET]
    ensure_tag_columns(ws)
    i, written = 0, 0
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if _cell(ws, r, "date") is None or not _cell(ws, r, "asset") \
           or _cell(ws, r, "entry") is None:
            continue
        label = labels.get(i)
        if isinstance(label, (list, tuple, set)):
            label = ", ".join(dict.fromkeys(label))  # de-dup, keep order
        cur = ws.cell(row=r, column=COL[column]).value
        cur_s = str(cur).strip() if cur else ""
        replaceable = bool(replace_values) and cur_s in replace_values
        if label and (not only_blank or not cur_s or replaceable):
            ws.cell(row=r, column=COL[column], value=label)
            written += 1
        i += 1
    wb.save(wb_path)
    wb.close()
    return written


def ensure_tag_columns(ws) -> None:
    """Make sure every tag-category column has its correct header."""
    from . import tags as _t
    for key, header in _t.FIELD_HEADER.items():
        ws.cell(row=1, column=COL[key], value=header)   # keep headers in sync


TAG_OPTIONS_SHEET = "Tag Options"


def ensure_tag_options_sheet(workbook: Path | None = None) -> None:
    """Create the 'Tag Options' sheet with defaults ONLY if it doesn't exist.

    Never overwrites an existing sheet — the trader owns its contents, so their
    added/edited options persist across runs.
    """
    from . import tags as _t
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    if TAG_OPTIONS_SHEET in wb.sheetnames:
        wb.close()
        return
    ws = wb.create_sheet(TAG_OPTIONS_SHEET)
    cats = list(_t.TAXONOMY.keys())
    for c, cat in enumerate(cats, start=1):
        hc = ws.cell(row=1, column=c, value=cat)
        hc.fill = _HDR_FILL
        hc.font = _HDR_FONT
        ws.column_dimensions[get_column_letter(c)].width = max(16, len(cat) + 2)
        for r, val in enumerate(_t.TAXONOMY[cat], start=2):
            ws.cell(row=r, column=c, value=val)
    ws.cell(row=1, column=len(cats) + 2,
            value="Suggested tag values — add/edit your own; they show as chips.").font = \
        Font(italic=True, color="6B7280")
    wb.save(wb_path)
    wb.close()


RULES_SHEET = "Trading Rules"


def ensure_rules_sheet(workbook: Path | None = None) -> None:
    """Create the 'Trading Rules' sheet with defaults ONLY if it doesn't exist."""
    from . import tags as _t
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    if RULES_SHEET in wb.sheetnames:
        wb.close()
        return
    ws = wb.create_sheet(RULES_SHEET)
    for c, h in enumerate(["Type", "Rule"], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 100
    r = 2
    for rule in _t.DEFAULT_HARD_RULES:
        ws.cell(row=r, column=1, value="Hard")
        ws.cell(row=r, column=2, value=rule).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for rule in _t.DEFAULT_OTHER_RULES:
        ws.cell(row=r, column=1, value="Other")
        ws.cell(row=r, column=2, value=rule).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    wb.save(wb_path)
    wb.close()


def read_rules(workbook: Path | None = None) -> dict:
    """Return {'hard': [...], 'other': [...]} from the Trading Rules sheet."""
    from . import tags as _t
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    try:
        wb = openpyxl.load_workbook(wb_path, data_only=True)
    except Exception:
        return {"hard": list(_t.DEFAULT_HARD_RULES), "other": list(_t.DEFAULT_OTHER_RULES)}
    if RULES_SHEET not in wb.sheetnames:
        wb.close()
        return {"hard": list(_t.DEFAULT_HARD_RULES), "other": list(_t.DEFAULT_OTHER_RULES)}
    ws = wb[RULES_SHEET]
    out = {"hard": [], "other": []}
    for r in range(2, ws.max_row + 1):
        typ = str(ws.cell(row=r, column=1).value or "").strip().lower()
        rule = ws.cell(row=r, column=2).value
        if not rule or not str(rule).strip():
            continue
        (out["hard"] if typ == "hard" else out["other"]).append(str(rule).strip())
    wb.close()
    return out


def write_rules(hard: list, other: list, workbook: Path | None = None) -> None:
    """Rewrite the Trading Rules sheet from the given lists."""
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    backup(wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    if RULES_SHEET in wb.sheetnames:
        del wb[RULES_SHEET]
    ws = wb.create_sheet(RULES_SHEET)
    for c, h in enumerate(["Type", "Rule"], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 100
    r = 2
    for typ, rules in (("Hard", hard), ("Other", other)):
        for rule in rules:
            if not str(rule).strip():
                continue
            ws.cell(row=r, column=1, value=typ)
            ws.cell(row=r, column=2, value=str(rule)).alignment = \
                Alignment(wrap_text=True, vertical="top")
            r += 1
    wb.save(wb_path)
    wb.close()


NEWS_SHEET = "Economic News"


def write_news_sheet(events: list[dict], workbook: Path | None = None) -> int:
    """Create/replace an 'Economic News' sheet of high-impact events.

    Derived data (regenerated each run), for reference and AI analysis.
    Returns the number of events written.
    """
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    if NEWS_SHEET in wb.sheetnames:
        del wb[NEWS_SHEET]
    ws = wb.create_sheet(NEWS_SHEET)
    cols = ["Date", "Time (ET)", "Currency", "Impact", "Event",
            "Actual", "Forecast", "Previous"]
    widths = [12, 10, 9, 8, 34, 10, 10, 10]
    for c, (h, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, e in enumerate(events, start=2):
        for c, key in enumerate(
                ["date", "time", "ccy", "impact", "event",
                 "actual", "forecast", "previous"], start=1):
            ws.cell(row=r, column=c, value=e.get(key, ""))
    wb.save(wb_path)
    wb.close()
    return len(events)


INDICATORS_SHEET = "Indicators"


def write_indicators_sheet(rows: list[dict], workbook: Path | None = None) -> int:
    """Create/replace an 'Indicators' sheet — one row per trade with its
    multi-timeframe indicator state at entry (derived, regenerated each run)."""
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    if INDICATORS_SHEET in wb.sheetnames:
        del wb[INDICATORS_SHEET]
    ws = wb.create_sheet(INDICATORS_SHEET)
    cols = [("Date", "date"), ("Time", "time"), ("Asset", "asset"),
            ("Dir", "dir"), ("P&L", "pnl"), ("1H Trend", "h1_trend"),
            ("1H ST", "h1_st"), ("15m ST", "m15_st"), ("5m ST", "m5_st"),
            ("5m RSI", "m5_rsi"), ("5m EMA", "m5_ema"), ("5m MACD", "m5_macd"),
            ("5m ADX", "m5_adx"), ("5m ATR", "m5_atr"),
            ("VWAP", "vwap"), ("VWAP Side", "vwap_side"), ("VPOC", "vpoc"),
            ("RVOL", "rvol"), ("Est. Delta", "est_delta"),
            ("Fib %", "fib_retr"), ("Fib Zone", "fib_zone"),
            ("Pullback", "pb_grade"), ("PB Score", "pb_score"),
            ("At Level", "at_level"), ("Candle", "candle"), ("Trig Vol", "trigvol"),
            ("STs Aligned", "aligned"), ("With 1H Trend", "with_1h")]
    for c, (h, _) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        ws.column_dimensions[get_column_letter(c)].width = 11
    for r, row in enumerate(rows, start=2):
        for c, (_, key) in enumerate(cols, start=1):
            ws.cell(row=r, column=c, value=row.get(key, ""))
    wb.save(wb_path)
    wb.close()
    return len(rows)


def read_tag_options(workbook: Path | None = None) -> dict:
    """Read the tag vocabulary from the 'Tag Options' sheet.

    Returns {category: [values]} for the known categories, using the sheet's
    values where a matching column exists, else the code defaults. Matching is
    by column header vs. category name (case-insensitive).
    """
    from . import tags as _t
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    try:
        wb = openpyxl.load_workbook(wb_path, data_only=True)
    except Exception:
        return dict(_t.TAXONOMY)
    if TAG_OPTIONS_SHEET not in wb.sheetnames:
        wb.close()
        return dict(_t.TAXONOMY)
    ws = wb[TAG_OPTIONS_SHEET]
    header_col = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if not h:
            continue
        for cat in _t.TAXONOMY:
            if str(h).strip().lower() == cat.lower():
                header_col[cat] = c
    vocab = {}
    for cat in _t.TAXONOMY:                    # keep canonical order
        col = header_col.get(cat)
        if col is None:
            vocab[cat] = list(_t.TAXONOMY[cat])
            continue
        vals = []
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v is not None and str(v).strip():
                vals.append(str(v).strip())
        vocab[cat] = vals
    wb.close()
    return vocab


def _row_matches(ws, r, ident: dict) -> bool:
    d = _cell(ws, r, "date")
    if d is None:
        return False
    d = d.date() if isinstance(d, datetime) else d
    if str(d) != str(ident.get("date")):
        return False
    tcell = _cell(ws, r, "time")
    tstr = tcell.strftime("%H:%M") if isinstance(tcell, time) else ""
    if ident.get("time") and tstr != ident["time"]:
        return False
    if str(_cell(ws, r, "asset") or "").upper() != str(ident.get("asset", "")).upper():
        return False
    entry = _num(_cell(ws, r, "entry"))
    ie = ident.get("entry")
    if entry is not None and ie is not None and round(entry, 4) != round(float(ie), 4):
        return False
    return True


def update_tags(ident: dict, values: dict, screenshot=None,
                workbook: Path | None = None) -> bool:
    """Write tag categories for the trade matching `ident`.

    `ident` keys: date (YYYY-MM-DD), time (HH:MM), asset, direction, entry.
    `values` maps field keys (setup, mistakes, rules, market, emotion, process,
    learning) to comma-separated strings; only provided keys are written.
    `screenshot` writes the Trade Screenshot column when not None.
    Returns True if a row matched and was written.
    """
    from . import tags as _t
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    backup(wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws = wb[config.TRADE_SHEET]
    ensure_tag_columns(ws)
    hit = False
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if _row_matches(ws, r, ident):
            for key in _t.FIELD_KEYS:
                if key in values:
                    ws.cell(row=r, column=COL[key], value=values[key] or "")
            if screenshot is not None:
                ws.cell(row=r, column=COL["trade_screenshot"], value=screenshot or "")
            hit = True
            break
    if hit:
        wb.save(wb_path)
    wb.close()
    return hit


def ensure_multipliers(ws_mult) -> None:
    """Make sure every configured root has a row on the Multipliers sheet."""
    existing = {}
    for r in range(2, ws_mult.max_row + 1):
        name = ws_mult.cell(row=r, column=1).value
        if name:
            existing[str(name).strip().upper()] = r
    next_row = ws_mult.max_row + 1
    for root, mult in config.MULTIPLIERS.items():
        if root not in existing:
            ws_mult.cell(row=next_row, column=1, value=root)
            ws_mult.cell(row=next_row, column=2, value=mult)
            next_row += 1


def _row_formulas(r: int) -> dict[str, str]:
    """The workbook's per-row formulas, with the Cumulative P&L bug fixed."""
    return {
        "rr_targeted": f'=IF(AND(I{r}<>"",J{r}<>"",F{r}<>""), (J{r}-F{r})/(F{r}-I{r}), "")',
        "rr_realized": f'=IF(AND(F{r}<>"",G{r}<>"",I{r}<>""), (G{r}-F{r})/(F{r}-I{r}), "")',
        "winloss": f'=IF(N{r}>0,"Win",IF(N{r}<0,"Loss",""))',
        "pnl": (f'=IF(AND(F{r}<>"",G{r}<>"",H{r}<>"",D{r}<>""), '
                f'IF(E{r}="Long", (G{r}-F{r})*H{r}*VLOOKUP(D{r},Multipliers!A:B,2,FALSE), '
                f'(F{r}-G{r})*H{r}*VLOOKUP(D{r},Multipliers!A:B,2,FALSE)), "")'),
        "pnl_pct": (f'=IF(AND(N{r}<>"",F{r}<>"",H{r}<>"",D{r}<>""), '
                    f'N{r}/(F{r}*H{r}*VLOOKUP(D{r},Multipliers!A:B,2,FALSE)), "")'),
        # Fixed: previous row reference instead of the broken #REF!.
        "cum_pnl": (f"=N{r}" if r == FIRST_DATA_ROW else f"=P{r-1}+N{r}"),
    }


def _last_real_row(ws) -> int:
    last = FIRST_DATA_ROW - 1
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if _cell(ws, r, "date") is not None and _cell(ws, r, "asset"):
            last = r
    return last


def append_trades(new_trades: list[Trade], workbook: Path | None = None) -> int:
    """Append new trades, writing raw inputs + the workbook's formulas.

    Also repairs the Cumulative P&L formula across all existing data rows.
    Returns the number of rows written.
    """
    if not new_trades:
        return 0
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    backup(wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws = wb[config.TRADE_SHEET]
    if config.MULTIPLIER_SHEET in wb.sheetnames:
        ensure_multipliers(wb[config.MULTIPLIER_SHEET])

    last = _last_real_row(ws)
    last_zth = _cell(ws, last, "zth") if last >= FIRST_DATA_ROW else None
    try:
        next_zth = int(last_zth) + 1
    except (TypeError, ValueError):
        next_zth = (last - FIRST_DATA_ROW + 2)

    r = last + 1
    for t in new_trades:
        ws.cell(row=r, column=COL["zth"], value=next_zth)
        d = t.date.date() if isinstance(t.date, datetime) else t.date
        c = ws.cell(row=r, column=COL["date"], value=datetime.combine(d, time(0, 0)) if d else None)
        c.number_format = "m/d/yyyy"
        c = ws.cell(row=r, column=COL["time"], value=t.time)
        c.number_format = "h:mm"
        ws.cell(row=r, column=COL["asset"], value=t.asset)
        ws.cell(row=r, column=COL["direction"], value=t.direction)
        ws.cell(row=r, column=COL["entry"], value=t.entry)
        ws.cell(row=r, column=COL["exit"], value=t.exit)
        ws.cell(row=r, column=COL["size"], value=t.size)
        ws.cell(row=r, column=COL["stop"], value=t.stop)
        ws.cell(row=r, column=COL["target"], value=t.target)
        ws.cell(row=r, column=COL["setup"], value=t.setup or "")
        ws.cell(row=r, column=COL["notes"], value=t.notes or "")
        ws.cell(row=r, column=COL["instrument_level"], value=t.asset)
        for key, f in _row_formulas(r).items():
            ws.cell(row=r, column=COL[key], value=f)
        next_zth += 1
        r += 1

    # Repair cumulative formulas for every real row (fixes historic #REF!).
    for rr in range(FIRST_DATA_ROW, r):
        if _cell(ws, rr, "date") is not None:
            ws.cell(row=rr, column=COL["cum_pnl"],
                    value=_row_formulas(rr)["cum_pnl"])

    wb.save(wb_path)
    wb.close()
    return len(new_trades)


# --- Analysis write-back ---------------------------------------------------
_HDR_FILL = PatternFill("solid", fgColor="1F2937")
_HDR_FONT = Font(color="FFFFFF", bold=True)
_SEC_FONT = Font(bold=True, size=13, color="111827")


def write_analysis_sheets(metrics: dict, ai_text: str,
                          workbook: Path | None = None) -> None:
    """Write/replace 'Analysis' and 'AI Suggestions' sheets in the workbook."""
    wb_path = Path(workbook) if workbook else config.WORKBOOK
    backup(wb_path)
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    for name in ("Analysis", "AI Suggestions"):
        if name in wb.sheetnames:
            del wb[name]

    a = wb.create_sheet("Analysis")
    a.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        a.column_dimensions[col].width = 14
    row = 1

    def section(title):
        nonlocal row
        c = a.cell(row=row, column=1, value=title)
        c.font = _SEC_FONT
        row += 1

    def header(cols):
        nonlocal row
        for i, h in enumerate(cols, start=1):
            c = a.cell(row=row, column=i, value=h)
            c.fill, c.font = _HDR_FILL, _HDR_FONT
        row += 1

    def line(cols):
        nonlocal row
        for i, v in enumerate(cols, start=1):
            a.cell(row=row, column=i, value=v)
        row += 1

    s = metrics["summary"]
    section("Overview")
    line(["Generated", metrics.get("generated_at", "")])
    line(["Date range", metrics.get("date_range", "")])
    line(["Trades", s["n_trades"]])
    line(["Win rate", f'{s["win_rate"]*100:.1f}%'])
    line(["Total P&L", round(s["total_pnl"], 2)])
    line(["Profit factor", s["profit_factor"]])
    line(["Expectancy / trade", round(s["expectancy"], 2)])
    line(["Avg win", round(s["avg_win"], 2)])
    line(["Avg loss", round(s["avg_loss"], 2)])
    line(["Avg RR targeted", s["avg_rr_targeted"]])
    line(["Avg RR realized", s["avg_rr_realized"]])
    line(["Max drawdown", round(s["max_drawdown"], 2)])
    line(["Max win streak", s["max_win_streak"]])
    line(["Max loss streak", s["max_loss_streak"]])
    row += 1

    def table(title, key, label):
        nonlocal row
        rows = metrics.get(key, [])
        if not rows:
            return
        section(title)
        header([label, "Trades", "Win %", "Total P&L", "Expectancy"])
        for b in rows:
            line([b["name"], b["n"], f'{b["win_rate"]*100:.0f}%',
                  round(b["total_pnl"], 2), round(b["expectancy"], 2)])
        row += 1

    table("By Setup", "by_setup", "Setup")
    table("By Instrument", "by_instrument", "Instrument")
    table("By Direction", "by_direction", "Direction")
    table("By Hour", "by_hour", "Hour")
    table("By Day of Week", "by_dayofweek", "Day")

    flags = metrics.get("behavior_flags", [])
    if flags:
        section("Behavior Flags")
        for f in flags:
            line([f["type"], f["detail"]])

    # AI Suggestions sheet
    ai = wb.create_sheet("AI Suggestions")
    ai.column_dimensions["A"].width = 120
    ai.cell(row=1, column=1, value="Claude AI Coaching").font = _SEC_FONT
    rr = 3
    for para in (ai_text or "No AI output.").split("\n"):
        c = ai.cell(row=rr, column=1, value=para)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    wb.save(wb_path)
    wb.close()
